from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook

from app.errors import DomainError

MAX_ROWS = 500

COLUMNS = {
    "name": "name",
    "email": "email",
    "country": "country_code",
    "country_code": "country_code",
    "department": "department",
    "dept": "department",
    "designation": "designation",
    "title": "designation",
    "employee_code": "employee_code",
}


def decode_excel_base64(file: str) -> bytes:
    payload = file[file.index(",") + 1 :] if "," in file else file
    import base64

    return base64.b64decode(payload)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat") and not isinstance(value, str):
        try:
            return value.isoformat()[:10]
        except Exception:
            return str(value).strip()
    return str(value).strip()


def parse_employee_excel(buffer: bytes) -> list[dict]:
    try:
        book = load_workbook(BytesIO(buffer), data_only=True, read_only=True)
    except Exception:
        raise DomainError("file must be a valid Excel workbook (.xlsx)") from None
    sheet = book.worksheets[0] if book.worksheets else None
    if sheet is None:
        raise DomainError("Excel file needs a worksheet")

    extracted: list[tuple[int, list[str]]] = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        cells = [_cell_text(cell) for cell in row]
        if all(not cell for cell in cells):
            continue
        extracted.append((row_number, cells))

    if not extracted or len(extracted) < 2:
        raise DomainError("Excel needs a header and at least one person")

    header = [COLUMNS.get(label.lower()) for label in extracted[0][1]]
    for required in ("name", "email", "country_code", "department", "designation"):
        if required not in header:
            raise DomainError(
                "Excel header must include name, email, country_code, department, designation",
            )

    def get(cells: list[str], key: str) -> str:
        try:
            return cells[header.index(key)]
        except ValueError:
            return ""

    people = []
    for line, cells in extracted[1:]:
        code = get(cells, "employee_code")
        people.append(
            {
                "name": get(cells, "name"),
                "email": get(cells, "email"),
                "country_code": get(cells, "country_code"),
                "department": get(cells, "department"),
                "designation": get(cells, "designation"),
                "employee_code": code or None,
                "line": line,
            }
        )
    if len(people) > MAX_ROWS:
        raise DomainError(f"Excel can have at most {MAX_ROWS} people")
    return people


def sample_employee_excel() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "People"
    sheet.append(["name", "email", "country_code", "department", "designation"])
    sheet.append(["Priya Shah", "excel.priya.shah@acme.test", "IN", "Engineering", "Software Engineer"])
    sheet.append(["James Chen", "excel.james.chen@acme.test", "US", "Product", "Product Manager"])
    out = BytesIO()
    book.save(out)
    return out.getvalue()
